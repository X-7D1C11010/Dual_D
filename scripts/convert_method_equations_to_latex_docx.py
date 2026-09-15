from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"D:\Code\Dual_D")
SOURCE = ROOT / "面向卫星多模态分布漂移的SAR与光学双对抗域适应网络_方法模块补充版.docx"
OUTPUT = ROOT / "面向卫星多模态分布漂移的SAR与光学双对抗域适应网络_方法模块LaTeX版.docx"


LATEX = {
    7: r"""\begin{aligned}
G_{a\rightarrow b}(f)&=f+\rho\tanh\!\left(g_{a\rightarrow b}(f)\right),\\
\hat{f}_{t}&=G_{s\rightarrow t}(f_s),\qquad
\hat{f}_{s}=G_{t\rightarrow s}(f_t).
\end{aligned}\tag{7}""",
    8: r"""\begin{aligned}
\mathcal{L}_{D}^{s}
&=\frac{1}{2}\left[\ell_{\mathrm{bin}}\!\left(D_s(f_s),1\right)
+\ell_{\mathrm{bin}}\!\left(D_s(\hat{f}_s),0\right)\right],\\
\mathcal{L}_{D}^{t}
&=\frac{1}{2}\left[\ell_{\mathrm{bin}}\!\left(D_t(f_t),1\right)
+\ell_{\mathrm{bin}}\!\left(D_t(\hat{f}_t),0\right)\right],\\
\mathcal{L}_{D}&=\frac{1}{2}\left(\mathcal{L}_{D}^{s}+\mathcal{L}_{D}^{t}\right).
\end{aligned}\tag{8}""",
    9: r"""\begin{aligned}
\mathcal{L}_{\mathrm{adv}}^{s}
&=\ell_{\mathrm{bin}}\!\left(D_s(\hat{f}_s),1\right),\qquad
\mathcal{L}_{\mathrm{adv}}^{t}
=\ell_{\mathrm{bin}}\!\left(D_t(\hat{f}_t),1\right),\\
\mathcal{L}_{\mathrm{adv}}
&=\lambda_{\mathrm{adv}}^{s}\mathcal{L}_{\mathrm{adv}}^{s}
+\lambda_{\mathrm{adv}}^{t}\mathcal{L}_{\mathrm{adv}}^{t}.
\end{aligned}\tag{9}""",
    10: r"""\begin{aligned}
\mathcal{L}_{\mathrm{cyc}}
={}&\left\|G_{t\rightarrow s}\!\left(G_{s\rightarrow t}(f_s)\right)-f_s\right\|_1\\
&+\left\|G_{s\rightarrow t}\!\left(G_{t\rightarrow s}(f_t)\right)-f_t\right\|_1 .
\end{aligned}\tag{10}""",
    11: r"""\mathcal{L}_{\mathrm{id}}
=\left\|G_{t\rightarrow s}(f_s)-f_s\right\|_1
+\left\|G_{s\rightarrow t}(f_t)-f_t\right\|_1 .\tag{11}""",
    12: r"""\ell_{\mathrm{mp}}(Q,R,a,b)
=-\frac{1}{|\mathcal{I}|}\sum_{i\in\mathcal{I}}
\log\frac{
\displaystyle\sum_{j:b_j=a_i}
\exp\!\left(\operatorname{sim}\!\left(q_i,\operatorname{sg}(r_j)\right)/\tau\right)}{
\displaystyle\sum_j
\exp\!\left(\operatorname{sim}\!\left(q_i,\operatorname{sg}(r_j)\right)/\tau\right)} .\tag{12}""",
    13: r"""\mathcal{L}_{\mathrm{mp}}
=\frac{1}{2}\left[
\ell_{\mathrm{mp}}(\hat{F}_s,F_s,y_t,y_s)
+\ell_{\mathrm{mp}}(\hat{F}_t,F_t,y_s,y_t)
\right].\tag{13}""",
    14: r"""\begin{aligned}
\bar{p}_{d,c}
&=\frac{1}{|\mathcal{B}_{d,c}|}
\sum_{i:y_{d,i}=c}\operatorname{sg}(f_{d,i}),\\
p_{d,c}
&\leftarrow\mu p_{d,c}+(1-\mu)\bar{p}_{d,c},
\qquad d\in\{s,t\}.
\end{aligned}\tag{14}""",
    15: r"""\begin{aligned}
\ell_{\mathrm{proto}}(Q,a,P_d)
&=-\frac{1}{|\mathcal{I}_d|}\sum_{i\in\mathcal{I}_d}
\log\frac{
\exp\!\left(\operatorname{sim}(q_i,p_{d,a_i})/\tau_p\right)}{
\displaystyle\sum_{c\in\mathcal{C}_d}
\exp\!\left(\operatorname{sim}(q_i,p_{d,c})/\tau_p\right)},\\
\mathcal{L}_{\mathrm{proto}}
&=\frac{1}{2}\left[
\ell_{\mathrm{proto}}(\hat{F}_s,y_t,P_s)
+\ell_{\mathrm{proto}}(\hat{F}_t,y_s,P_t)
\right].
\end{aligned}\tag{15}""",
    16: r"""\mathcal{L}_{\mathrm{fb}}
=\operatorname{CE}\!\left(C(\hat{f}_s),y_t\right)
+\operatorname{CE}\!\left(C(\hat{f}_t),y_s\right).\tag{16}""",
    17: r"""\mathcal{L}_{\mathrm{cls}}
=\operatorname{CE}\!\left(C(f_s),y_s\right)
+\lambda_{\mathrm{cls}}^{t}
\operatorname{CE}\!\left(C(f_t),y_t\right).\tag{17}""",
    18: r"""\mathcal{L}_{C}
=\lambda_{\mathrm{cyc}}\mathcal{L}_{\mathrm{cyc}}
+\lambda_{\mathrm{id}}\mathcal{L}_{\mathrm{id}}
+\lambda_{\mathrm{mp}}\mathcal{L}_{\mathrm{mp}}
+\lambda_{\mathrm{proto}}\mathcal{L}_{\mathrm{proto}}
+\lambda_{\mathrm{fb}}\mathcal{L}_{\mathrm{fb}}.\tag{18}""",
    19: r"""\begin{aligned}
\bar{Z}_{i}^{(m)}
&=\frac{Z_{i}^{(m)}}{\left\|Z_{i}^{(m)}\right\|_2+\varepsilon},\qquad
R^{(m)}=\bar{Z}^{(m)}\bar{Z}^{(m)\top},\\
A(F)&=\frac{1}{B^2}
\left\|R^{(\mathrm{sar})}-R^{(\mathrm{opt})}\right\|_F^2 .
\end{aligned}\tag{19}""",
    20: r"""\begin{aligned}
\Delta_{s\rightarrow t}
&=A\!\left(G_{s\rightarrow t}\!\left(\operatorname{sg}(F_s)\right)\right)
-\operatorname{sg}\!\left(A(F_s)\right),\\
\Delta_{t\rightarrow s}
&=A\!\left(G_{t\rightarrow s}\!\left(\operatorname{sg}(F_t)\right)\right)
-\operatorname{sg}\!\left(A(F_t)\right).
\end{aligned}\tag{20}""",
    21: r"""\mathcal{L}_{\mathrm{drift}}
=\frac{1}{2}\left(
[\Delta_{s\rightarrow t}-\delta]_+
+[\Delta_{t\rightarrow s}-\delta]_+
\right),\qquad [u]_+=\max(u,0).\tag{21}""",
    22: r"""\mathcal{L}_{\mathrm{main}}
=\mathcal{L}_{\mathrm{cls}}
+\lambda_{\mathrm{tal}}\mathcal{L}_{\mathrm{tal}}
+\gamma_{\mathrm{adv}}\mathcal{L}_{\mathrm{adv}}
+\gamma_C\mathcal{L}_{C}
+\gamma_{\mathrm{drift}}\lambda_{\mathrm{drift}}
\mathcal{L}_{\mathrm{drift}}.\tag{22}""",
}


def paragraph_all_text(paragraph) -> str:
    return "".join(paragraph._p.itertext())


def replace_paragraph_with_latex(paragraph, latex: str) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)

    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(0)
    fmt.left_indent = Pt(21)
    fmt.right_indent = Pt(10)
    fmt.line_spacing = Pt(12)
    fmt.space_before = Pt(3)
    fmt.space_after = Pt(3)
    fmt.keep_together = True

    run = paragraph.add_run(latex)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")


def replace_text(document, old: str, new: str) -> None:
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            full = paragraph.text.replace(old, new)
            paragraph.clear()
            run = paragraph.add_run(full)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def main() -> None:
    document = Document(SOURCE)

    paragraphs = document.paragraphs
    start = next(
        i for i, paragraph in enumerate(paragraphs)
        if paragraph.text.strip() == "双生成器双判别器模块"
    )
    end = next(
        i for i, paragraph in enumerate(paragraphs[start + 1 :], start + 1)
        if paragraph.text.strip() == "实验设置"
    )

    found = set()
    for paragraph in paragraphs[start:end]:
        text = paragraph_all_text(paragraph)
        for number, latex in LATEX.items():
            if f"({number})" in text:
                replace_paragraph_with_latex(paragraph, latex)
                found.add(number)
                break

    missing = set(LATEX) - found
    if missing:
        raise RuntimeError(f"未找到待替换公式：{sorted(missing)}")

    replace_text(
        document,
        "为此，本文进一步构造类感知的多正样本对比损失。设Q={q_i}为生成锚点，R={r_j}为对应输出域的真实候选，a和b为二者的类别标签，",
        "为此，本文进一步构造类感知的多正样本对比损失。对于大小为B的批次，记F_d=[f_{d,1},…,f_{d,B}]^T和F̂_d=[f̂_{d,1},…,f̂_{d,B}]^T分别为域d的真实特征矩阵与生成特征矩阵，y_d为对应的标签向量。设Q={q_i}为生成锚点，R={r_j}为输出域的真实候选，a和b为二者的类别标签，",
    )
    replace_text(document, "λ_TAL=0.3", "λ_tal=0.3")

    properties = document.core_properties
    properties.title = "面向卫星多模态分布漂移的SAR与光学双对抗域适应网络 方法模块LaTeX版"
    properties.subject = "方法模块公式LaTeX源码与符号统一"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
