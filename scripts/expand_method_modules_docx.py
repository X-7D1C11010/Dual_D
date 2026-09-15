from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"D:\Code\Dual_D")
SOURCE = ROOT / "面向卫星多模态分布漂移的SAR与光学双对抗域适应网络_M4-SAR修订版.docx"
OUTPUT = ROOT / "面向卫星多模态分布漂移的SAR与光学双对抗域适应网络_方法模块补充版.docx"
EQUATIONS = {
    "eq06": "Gₐ→ᵦ(f)=f+ρ tanh(gₐ→ᵦ(f));    f̂ₜ=Gₛ→ₜ(fₛ),    f̂ₛ=Gₜ→ₛ(fₜ).                                      (7)",
    "eq07": "ℒᴰˢ=½[ℓᵦᵢₙ(Dₛ(fₛ),1)+ℓᵦᵢₙ(Dₛ(f̂ₛ),0)];    ℒᴰᵗ=½[ℓᵦᵢₙ(Dₜ(fₜ),1)+ℓᵦᵢₙ(Dₜ(f̂ₜ),0)];    ℒᴰ=½(ℒᴰˢ+ℒᴰᵗ).    (8)",
    "eq08": "ℒₐᵈᵛˢ=ℓᵦᵢₙ(Dₛ(f̂ₛ),1),    ℒₐᵈᵛᵗ=ℓᵦᵢₙ(Dₜ(f̂ₜ),1);    ℒₐᵈᵛ=λₐᵈᵛˢℒₐᵈᵛˢ+λₐᵈᵛᵗℒₐᵈᵛᵗ.          (9)",
    "eq09": "ℒcyc=‖Gₜ→ₛ(Gₛ→ₜ(fₛ))−fₛ‖₁+‖Gₛ→ₜ(Gₜ→ₛ(fₜ))−fₜ‖₁.                                  (10)",
    "eq10": "ℒid=‖Gₜ→ₛ(fₛ)−fₛ‖₁+‖Gₛ→ₜ(fₜ)−fₜ‖₁.                                                          (11)",
    "eq11": "ℓmp(Q,R,a,b)=−|ℐ|⁻¹ Σᵢ∈ℐ log{Σⱼ:bⱼ=aᵢ exp[sim(qᵢ,sg(rⱼ))/τ] / Σⱼ exp[sim(qᵢ,sg(rⱼ))/τ]}.    (12)",
    "eq12": "ℒmp=½[ℓmp(F̂ₛ,Fₛ,yₜ,yₛ)+ℓmp(F̂ₜ,Fₜ,yₛ,yₜ)].                                               (13)",
    "eq13": "p̄d,c=|ℬd,c|⁻¹ Σᵢ:yd,i=c sg(fd,i);    pd,c←μpd,c+(1−μ)p̄d,c,    d∈{s,t}.                        (14)",
    "eq14": "ℓproto(Q,a,Pd)=−|ℐd|⁻¹Σᵢ∈ℐd log{exp[sim(qᵢ,pd,aᵢ)/τp] / Σc∈𝒞d exp[sim(qᵢ,pd,c)/τp]};    ℒproto=½[ℓproto(F̂ₛ,yₜ,Pₛ)+ℓproto(F̂ₜ,yₛ,Pₜ)].    (15)",
    "eq15": "ℒfb=CE(C(f̂ₛ),yₜ)+CE(C(f̂ₜ),yₛ).                                                               (16)",
    "eq16": "ℒcls=CE(C(fₛ),yₛ)+λclsᵗ CE(C(fₜ),yₜ).                                                         (17)",
    "eq17": "ℒC=λcycℒcyc+λidℒid+λmpℒmp+λprotoℒproto+λfbℒfb.                                             (18)",
    "eq18": "Z̄ᵢ⁽ᵐ⁾=Zᵢ⁽ᵐ⁾/(‖Zᵢ⁽ᵐ⁾‖₂+ε),    R⁽ᵐ⁾=Z̄⁽ᵐ⁾Z̄⁽ᵐ⁾ᵀ;    A(F)=B⁻²‖R⁽ˢᵃʳ⁾−R⁽ᵒᵖᵗ⁾‖F².      (19)",
    "eq19": "Δₛ→ₜ=A(Gₛ→ₜ(sg(Fₛ)))−sg(A(Fₛ));    Δₜ→ₛ=A(Gₜ→ₛ(sg(Fₜ)))−sg(A(Fₜ)).                         (20)",
    "eq20": "ℒdrift=½([Δₛ→ₜ−δ]₊+[Δₜ→ₛ−δ]₊),    [u]₊=max(u,0).                                           (21)",
    "eq21": "ℒmain=ℒcls+λTALℒTAL+γadvℒadv+γCℒC+γdriftλdriftℒdrift.                                      (22)",
}


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=10.5):
    run.font.name = latin
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def format_body(paragraph):
    pf = paragraph.paragraph_format
    pf.first_line_indent = Pt(21)
    pf.line_spacing = Pt(16)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.widow_control = True
    for run in paragraph.runs:
        set_run_font(run)


def format_heading(paragraph):
    pf = paragraph.paragraph_format
    pf.keep_with_next = True
    pf.space_before = Pt(4)
    pf.space_after = Pt(2)
    for run in paragraph.runs:
        set_run_font(run, east_asia="黑体", size=10.5)
        run.bold = True


def format_equation(paragraph):
    pf = paragraph.paragraph_format
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    pf.keep_together = True


def find_heading(document, text):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise RuntimeError(f"未找到标题：{text}")


def remove_between(start, end):
    node = start._p.getnext()
    parent = start._p.getparent()
    while node is not None and node is not end._p:
        following = node.getnext()
        parent.remove(node)
        node = following


def insert_text_before(document, anchor, text):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.add_run(text)
    anchor._p.addprevious(paragraph._p)
    format_body(paragraph)
    return paragraph


def insert_equation_before(document, anchor, name):
    paragraph = document.add_paragraph(style="Equation")
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_props = OxmlElement("m:rPr")
    math_style = OxmlElement("m:sty")
    math_style.set(qn("m:val"), "p")
    math_props.append(math_style)
    math_text = OxmlElement("m:t")
    math_text.text = EQUATIONS[name]
    math_run.append(math_props)
    math_run.append(math_text)
    omath.append(math_run)
    omath_para.append(omath)
    paragraph._p.append(omath_para)
    anchor._p.addprevious(paragraph._p)
    format_equation(paragraph)
    return paragraph


def insert_blocks(document, anchor, blocks):
    for kind, value in blocks:
        if kind == "text":
            insert_text_before(document, anchor, value)
        elif kind == "eq":
            insert_equation_before(document, anchor, value)
        else:
            raise ValueError(kind)


def main():
    document = Document(SOURCE)
    if "Equation" not in [style.name for style in document.styles]:
        document.styles.add_style("Equation", 1)

    double = find_heading(document, "双生成器双判别器")
    category = find_heading(document, "类别感知语义反馈")
    drift = find_heading(document, "模态关系漂移约束")
    overall = find_heading(document, "总体目标与优化")
    experiment = find_heading(document, "实验设置")

    double.text = "双生成器双判别器模块"
    category.text = "类别感知反馈模块"
    for heading in (double, category, drift, overall):
        format_heading(heading)

    remove_between(double, category)
    remove_between(category, drift)
    remove_between(drift, overall)
    remove_between(overall, experiment)

    insert_blocks(document, category, [
        ("text", "张量对齐通过线性投影建立跨模态、跨领域的相关结构，但地区形态、观测几何和传感器响应变化还会造成复杂的非线性偏移。为补偿这部分差异，本文在融合特征空间中设置源域到目标域生成器G_{s→t}和目标域到源域生成器G_{t→s}。两个生成器均采用残差多层感知机，将输入特征保留为主分支，只对领域差异学习有界残差："),
        ("eq", "eq06"),
        ("text", "其中，g_{a→b}表示残差映射，ρ控制残差幅度；f_s和f_t分别为张量模块输出的源域与目标域融合特征，f̂_t和f̂_s分别表示目标域化特征与源域化特征。生成器作用于完整融合向量，可以同时调整两个模态子空间并建立跨子空间联系，因此其作用对象是融合特征，而不是彼此独立的SAR或光学特征。"),
        ("text", "为分别评价两个转换方向，网络配置源域判别器D_s和目标域判别器D_t。D_s区分真实源域特征f_s与目标域样本生成的源域化特征f̂_s，D_t区分真实目标域特征f_t与源域样本生成的目标域化特征f̂_t。令ℓ_bin表示二类交叉熵，两个判别器的训练目标为："),
        ("eq", "eq07"),
        ("text", "更新判别器时，真实特征与生成特征均从主网络中分离，只优化D_s和D_t，使判别器先获得稳定的领域判断能力。随后固定判别器参数，利用其输出反向约束生成器："),
        ("eq", "eq08"),
        ("text", "式（9）分别推动f̂_s接近真实源域分布、f̂_t接近真实目标域分布。两个方向的判别信号既避免将源域与目标域简单压入同一混合分布，也为后续循环重建和双向类别反馈提供对应的生成特征。需要说明的是，对抗目标只约束特征的领域属性，不能单独保证转换前后的类别语义不变，因此还需引入类别感知反馈。"),
    ])

    insert_blocks(document, drift, [
        ("text", "领域判别器只判断生成特征是否符合目标领域的统计特征，并不检查转换前后的目标类别是否一致。为避免生成器以改变类别语义的方式降低对抗损失，本文从特征内容、批内样本关系、跨批类别中心和分类决策四个层面建立反馈。首先，利用双向循环一致性限制往返转换造成的信息丢失："),
        ("eq", "eq09"),
        ("text", "循环一致性要求源域特征经“源域—目标域—源域”转换后能够恢复，目标域特征亦然。在此基础上，恒等保持约束生成器对已经位于输出领域的真实特征不作无意义改写："),
        ("eq", "eq10"),
        ("text", "循环一致性和恒等保持主要约束特征内容，但不能直接规定同类样本在转换后的相对位置。为此，本文进一步构造类感知的多正样本对比损失。设Q={q_i}为生成锚点，R={r_j}为对应输出域的真实候选，a和b为二者的类别标签，sim(·,·)为余弦相似度，τ为温度系数，I表示在候选集合中至少存在一个同类样本的有效锚点集合，则："),
        ("eq", "eq11"),
        ("text", "式（12）将候选域中的全部同类样本共同视为正样本，将异类样本作为负样本；真实候选停止梯度，使该项以真实领域结构为参照约束生成路径。两个转换方向分别以源域和目标域真实特征作为候选，并取平均得到："),
        ("eq", "eq12"),
        ("text", "批内对比能够利用当前批次的细粒度类别关系，但其参照会随批次组成变化。为获得跨批次的稳定类别中心，本文分别维护源域和目标域的指数滑动原型。设B_{d,c}为域d中类别c的当前批次样本集合，p̄_{d,c}为该类的批均值，p_{d,c}为历史原型，则："),
        ("eq", "eq13"),
        ("text", "首次出现的类别用当前批均值初始化；尚未建立原型的类别不参与计算。源域化特征与源域原型比较，目标域化特征与目标域原型比较，原型对比损失为："),
        ("eq", "eq14"),
        ("text", "其中，C_d为域d中已经建立原型的类别集合，I_d只保留标签对应原型有效的锚点。样本级对比约束局部类结构，原型对比提供跨批次的类别尺度参照，两者共同减少同类特征在转换后分散或偏向其他类别的问题。"),
        ("text", "最后，将生成特征送入共享分类器C，以原始输入样本的标签检查转换后的类别可判别性："),
        ("eq", "eq15"),
        ("text", "计算式（16）时冻结分类器参数，梯度通过固定的分类边界传回生成路径，使生成特征向其原类别的判别区域调整，而不是让分类器迁就错误转换。与此同时，真实源域与目标域融合特征用于训练编码器、张量投影和分类器："),
        ("eq", "eq16"),
        ("text", "综合上述约束，类别感知反馈项写为："),
        ("eq", "eq17"),
        ("text", "由此，循环一致性和恒等保持负责保存特征内容，样本对比与类别原型维持局部和全局类别结构，分类器反馈则从决策边界检查生成结果。各项共同作用于双向转换过程，使领域特征改变的同时尽量保留目标类别语义。"),
    ])

    insert_blocks(document, overall, [
        ("text", "双向生成器对完整融合特征进行非线性调整，虽然能够缩小领域差异，但这种整体变换也可能改变SAR与光学子空间对样本关系的描述，使张量阶段形成的跨模态协同关系发生退化。模态关系漂移约束不承担新的领域对齐任务，而是以张量投影结果为参照，检查转换前后的模态关系是否出现明显恶化。"),
        ("text", "对于批量融合特征F，将其按张量投影的固定边界划分为SAR子空间Z^(sar)和光学子空间Z^(opt)。由于两个子空间的坐标含义不同，本文不直接最小化两类特征之间的欧氏距离，而是先对各样本进行行归一化，再分别构造批内余弦关系矩阵R^(sar)和R^(opt)。两种模态对样本关系描述的不一致度定义为："),
        ("eq", "eq18"),
        ("text", "A(F)越小，表示两个模态子空间对当前批次样本相似关系的描述越接近。该量只用于比较领域转换前后的相对变化，不作为模态已经完全对齐的判据。对源域到目标域和目标域到源域两个方向，分别计算关系不一致度的变化量："),
        ("eq", "eq19"),
        ("text", "其中，sg(·)表示停止梯度。转换前的张量投影结果作为固定基线；专门计算漂移项时，生成器输入同样停止梯度，使该分支只调整双向生成器，不推动编码器或张量投影改变参照。考虑到领域转换需要保留必要的特征变化，本文设置容忍边界δ，并仅惩罚超过该边界的关系恶化："),
        ("eq", "eq20"),
        ("text", "当转换后两种模态的样本关系更加一致，或不一致度的增加未超过δ时，漂移损失为零；只有关系恶化超出容忍范围时才产生梯度。训练初期先建立基本的领域转换能力，随后逐步启用该约束，以避免对必要的特征调整施加过早限制。这样，领域转换可以缩小源域与目标域差异，同时尽量保留张量阶段已经形成的跨模态关系。由于生成器允许跨子空间混合，转换后的两段特征更准确地说是由张量投影边界定义的SAR与光学子空间位置，而不是完全独立的单模态特征。"),
    ])

    insert_blocks(document, experiment, [
        ("text", "训练采用判别器与主网络交替更新。首先固定编码器、张量投影和生成器，使用式（8）更新D_s与D_t；随后固定两个判别器，联合优化真实特征分类、张量对齐、生成器对抗目标、类别感知反馈和模态关系漂移约束："),
        ("eq", "eq21"),
        ("text", "其中，γ_adv、γ_C和γ_drift控制三个辅助分支的分阶段启用。每次主网络更新后，对张量投影矩阵执行正交回缩。整个训练顺序形成连续闭环：张量模块先建立模态—领域关系，双向生成器和双判别器补偿非线性领域差异，类别感知反馈维持转换前后的类别语义，漂移约束则限制领域转换对已有模态关系的破坏。推理时只保留模态编码器、目标域张量投影、目标域到源域生成器和共享分类器；反向生成器、两个判别器及各辅助损失均不参与预测。"),
    ])

    properties = document.core_properties
    properties.title = "面向卫星多模态分布漂移的SAR与光学双对抗域适应网络（方法模块补充版）"
    properties.subject = "双生成器双判别器、类别感知反馈与模态关系漂移约束的流程和公式补充"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
