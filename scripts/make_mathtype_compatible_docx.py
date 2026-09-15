from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"D:\Code\Dual_D")
SOURCE = ROOT / "面向卫星多模态分布漂移的SAR与光学双对抗域适应网络_方法模块LaTeX版.docx"
OUTPUT = ROOT / "面向卫星多模态分布漂移的SAR与光学双对抗域适应网络_方法模块MathType兼容版.docx"


# MathType accepts a subset of TeX.  Keep each expression on one line and use
# only common mathematical commands; numbering is placed in the prose rather
# than inside the pasted expression.
MATHTYPE = {
    7: r"G_{a\rightarrow b}(f)=f+\rho\tanh(g_{a\rightarrow b}(f)),\quad \hat{f}_t=G_{s\rightarrow t}(f_s),\quad \hat{f}_s=G_{t\rightarrow s}(f_t)",
    8: r"\mathcal{L}_D^s={1\over 2}[\ell_{\mathrm{bin}}(D_s(f_s),1)+\ell_{\mathrm{bin}}(D_s(\hat{f}_s),0)],\quad \mathcal{L}_D^t={1\over 2}[\ell_{\mathrm{bin}}(D_t(f_t),1)+\ell_{\mathrm{bin}}(D_t(\hat{f}_t),0)],\quad \mathcal{L}_D={1\over 2}(\mathcal{L}_D^s+\mathcal{L}_D^t)",
    9: r"\mathcal{L}_{\mathrm{adv}}^s=\ell_{\mathrm{bin}}(D_s(\hat{f}_s),1),\quad \mathcal{L}_{\mathrm{adv}}^t=\ell_{\mathrm{bin}}(D_t(\hat{f}_t),1),\quad \mathcal{L}_{\mathrm{adv}}=\lambda_{\mathrm{adv}}^s\mathcal{L}_{\mathrm{adv}}^s+\lambda_{\mathrm{adv}}^t\mathcal{L}_{\mathrm{adv}}^t",
    10: r"\mathcal{L}_{\mathrm{cyc}}=\left\|G_{t\rightarrow s}(G_{s\rightarrow t}(f_s))-f_s\right\|_1+\left\|G_{s\rightarrow t}(G_{t\rightarrow s}(f_t))-f_t\right\|_1",
    11: r"\mathcal{L}_{\mathrm{id}}=\left\|G_{t\rightarrow s}(f_s)-f_s\right\|_1+\left\|G_{s\rightarrow t}(f_t)-f_t\right\|_1",
    12: r"\ell_{\mathrm{mp}}(Q,R,a,b)=-{1\over |\mathcal{I}|}\sum_{i\in\mathcal{I}}\log{\sum_{j:b_j=a_i}\exp(s(q_i,\mathrm{sg}(r_j))/\tau)\over \sum_j\exp(s(q_i,\mathrm{sg}(r_j))/\tau)}",
    13: r"\mathcal{L}_{\mathrm{mp}}={1\over 2}[\ell_{\mathrm{mp}}(\hat{F}_s,F_s,y_t,y_s)+\ell_{\mathrm{mp}}(\hat{F}_t,F_t,y_s,y_t)]",
    14: r"\bar{p}_{d,c}={1\over |\mathcal{B}_{d,c}|}\sum_{i:y_{d,i}=c}\mathrm{sg}(f_{d,i}),\quad p_{d,c}\leftarrow \mu p_{d,c}+(1-\mu)\bar{p}_{d,c},\quad d\in\{s,t\}",
    15: r"\ell_{\mathrm{proto}}(Q,a,P_d)=-{1\over |\mathcal{I}_d|}\sum_{i\in\mathcal{I}_d}\log{\exp(s(q_i,p_{d,a_i})/\tau_p)\over \sum_{c\in\mathcal{C}_d}\exp(s(q_i,p_{d,c})/\tau_p)},\quad \mathcal{L}_{\mathrm{proto}}={1\over 2}[\ell_{\mathrm{proto}}(\hat{F}_s,y_t,P_s)+\ell_{\mathrm{proto}}(\hat{F}_t,y_s,P_t)]",
    16: r"\mathcal{L}_{\mathrm{fb}}=\mathrm{CE}(C(\hat{f}_s),y_t)+\mathrm{CE}(C(\hat{f}_t),y_s)",
    17: r"\mathcal{L}_{\mathrm{cls}}=\mathrm{CE}(C(f_s),y_s)+\lambda_{\mathrm{cls}}^t\mathrm{CE}(C(f_t),y_t)",
    18: r"\mathcal{L}_C=\lambda_{\mathrm{cyc}}\mathcal{L}_{\mathrm{cyc}}+\lambda_{\mathrm{id}}\mathcal{L}_{\mathrm{id}}+\lambda_{\mathrm{mp}}\mathcal{L}_{\mathrm{mp}}+\lambda_{\mathrm{proto}}\mathcal{L}_{\mathrm{proto}}+\lambda_{\mathrm{fb}}\mathcal{L}_{\mathrm{fb}}",
    19: r"\bar{Z}_i^{(m)}={Z_i^{(m)}\over \left\|Z_i^{(m)}\right\|_2+\varepsilon},\quad R^{(m)}=\bar{Z}^{(m)}\bar{Z}^{(m)T},\quad A(F)={1\over B^2}\left\|R^{(\mathrm{sar})}-R^{(\mathrm{opt})}\right\|_F^2",
    20: r"\Delta_{s\rightarrow t}=A(G_{s\rightarrow t}(\mathrm{sg}(F_s)))-\mathrm{sg}(A(F_s)),\quad \Delta_{t\rightarrow s}=A(G_{t\rightarrow s}(\mathrm{sg}(F_t)))-\mathrm{sg}(A(F_t))",
    21: r"\mathcal{L}_{\mathrm{drift}}={1\over 2}([\Delta_{s\rightarrow t}-\delta]_+ + [\Delta_{t\rightarrow s}-\delta]_+),\quad [u]_+=\max(u,0)",
    22: r"\mathcal{L}_{\mathrm{main}}=\mathcal{L}_{\mathrm{cls}}+\lambda_{\mathrm{tal}}\mathcal{L}_{\mathrm{tal}}+\gamma_{\mathrm{adv}}\mathcal{L}_{\mathrm{adv}}+\gamma_C\mathcal{L}_C+\gamma_{\mathrm{drift}}\lambda_{\mathrm{drift}}\mathcal{L}_{\mathrm{drift}}",
}


PROSE_REPLACEMENTS = {
    "只对领域差异学习有界残差：": "只对领域差异学习有界残差，具体形式见式（7）：",
    "两个判别器的训练目标为：": "两个判别器的训练目标见式（8）：",
    "利用其输出反向约束生成器：": "利用其输出反向约束生成器，目标函数见式（9）：",
    "首先，利用双向循环一致性限制往返转换造成的信息丢失：": "首先，利用双向循环一致性限制往返转换造成的信息丢失，具体形式见式（10）：",
    "不作无意义改写：": "不作无意义改写，恒等保持损失见式（11）：",
    "有效锚点集合，则：": "有效锚点集合，多正样本对比损失见式（12）：",
    "并取平均得到：": "并取平均得到式（13）：",
    "历史原型，则：": "历史原型，更新规则见式（14）：",
    "原型对比损失为：": "原型对比损失见式（15）：",
    "检查转换后的类别可判别性：": "检查转换后的类别可判别性，分类反馈见式（16）：",
    "真实源域与目标域融合特征用于训练编码器、张量投影和分类器：": "真实源域与目标域融合特征用于训练编码器、张量投影和分类器，监督目标见式（17）：",
    "类别感知反馈项写为：": "类别感知反馈项汇总为式（18）：",
    "两种模态对样本关系描述的不一致度定义为：": "两种模态对样本关系描述的不一致度见式（19）：",
    "分别计算关系不一致度的变化量：": "分别计算关系不一致度的变化量，见式（20）：",
    "并仅惩罚超过该边界的关系恶化：": "并仅惩罚超过该边界的关系恶化，具体形式见式（21）：",
    "生成器对抗目标、类别感知反馈和模态关系漂移约束：": "生成器对抗目标、类别感知反馈和模态关系漂移约束，总体目标见式（22）：",
    "sim(·,·)为余弦相似度": "s(·,·)为余弦相似度",
}


def set_code_paragraph(paragraph, code: str) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(0)
    fmt.left_indent = Pt(21)
    fmt.right_indent = Pt(10)
    fmt.line_spacing = Pt(12)
    fmt.space_before = Pt(3)
    fmt.space_after = Pt(3)
    fmt.keep_together = False
    run = paragraph.add_run(f"${code}$")
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")


def replace_body_text(paragraph, old: str, new: str) -> bool:
    if old not in paragraph.text:
        return False
    text = paragraph.text.replace(old, new)
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    return True


def main() -> None:
    document = Document(SOURCE)
    paragraphs = document.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "双生成器双判别器模块")
    end = next(i for i, p in enumerate(paragraphs[start + 1 :], start + 1) if p.text.strip() == "实验设置")

    found = set()
    for paragraph in paragraphs[start:end]:
        for number, code in MATHTYPE.items():
            if f"\\tag{{{number}}}" in paragraph.text:
                set_code_paragraph(paragraph, code)
                found.add(number)
                break

    missing = set(MATHTYPE) - found
    if missing:
        raise RuntimeError(f"未找到待替换公式：{sorted(missing)}")

    replacement_hits = {old: 0 for old in PROSE_REPLACEMENTS}
    for paragraph in paragraphs[start:end]:
        for old, new in PROSE_REPLACEMENTS.items():
            if replace_body_text(paragraph, old, new):
                replacement_hits[old] += 1

    no_hit = [old for old, count in replacement_hits.items() if count == 0]
    if no_hit:
        raise RuntimeError(f"未找到待更新的公式引导语：{no_hit}")

    properties = document.core_properties
    properties.title = "面向卫星多模态分布漂移的SAR与光学双对抗域适应网络 方法模块MathType兼容版"
    properties.subject = "可直接粘贴至MathType的兼容TeX公式"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
