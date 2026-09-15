from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(r"D:\Code\Dual_D")
SOURCE = ROOT / "用于恶劣天气下异构光学与AIS数据的海上目标识别的双对抗网络.docx"
OUTPUT = ROOT / "面向卫星多模态分布漂移的SAR与光学双对抗域适应网络_最新修订稿.docx"


def set_run_font(run, *, east_asia="宋体", latin="Times New Roman", size=10.5, bold=None, italic=None):
    run.font.name = latin
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def body_format(paragraph):
    paragraph.paragraph_format.first_line_indent = Pt(21)
    paragraph.paragraph_format.line_spacing = Pt(16)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.widow_control = True
    for run in paragraph.runs:
        set_run_font(run)


def heading_format(paragraph, level):
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(7 if level == 1 else 4)
    paragraph.paragraph_format.space_after = Pt(2)
    for run in paragraph.runs:
        set_run_font(run, east_asia="黑体", size=12 if level == 1 else 10.5, bold=True)


def equation_format(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(6)
    paragraph.paragraph_format.right_indent = Pt(6)
    paragraph.paragraph_format.line_spacing = Pt(15)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_together = True
    for run in paragraph.runs:
        set_run_font(run, east_asia="Cambria Math", latin="Cambria Math", size=10, italic=True)


def title_format(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        set_run_font(run, east_asia="黑体", size=16, bold=True)


def add(document, text, style="Normal"):
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)
    if style == "Title":
        title_format(paragraph)
    elif style == "Heading 1":
        heading_format(paragraph, 1)
    elif style == "Heading 2":
        heading_format(paragraph, 2)
    elif style == "Equation":
        equation_format(paragraph)
    else:
        body_format(paragraph)
    return paragraph


def clear_body(document):
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(document):
    if "Equation" not in [style.name for style in document.styles]:
        document.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def build_document(document):
    add(document, "面向卫星多模态分布漂移的SAR与光学双对抗域适应网络", "Title")
    add(
        document,
        "摘要：卫星遥感数据的统计分布会随光照条件、季节物候、大气状态、观测几何、传感器响应和地理区域持续变化。"
        "这种漂移不仅改变单一模态的外观，还会改变合成孔径雷达与多光谱光学观测之间的互补关系，使在固定区域训练的分类模型难以稳定迁移。"
        "现有方法通常将多模态融合与领域对齐分开处理：前者侧重融合异构观测，后者侧重缩小源域与目标域差异，却很少约束后续领域变换是否破坏已经形成的模态协同结构。"
        "本文提出Satellite-Dual_D，一种面向SAR与光学卫星数据的双向对抗域适应网络。该方法首先使用参数独立的模态编码器提取SAR散射表征和光学光谱表征，"
        "再以因子化张量收缩建模跨模态高阶关系，并对齐源域与目标域中同一模态的上下文增强表示。随后，双向残差生成器与两个方向特定判别器学习源域到目标域、"
        "目标域到源域的非线性特征转换。循环一致性、恒等保持、同类多正样本对比、指数滑动类别原型和分类器反馈共同限制转换过程中的语义偏移。"
        "最后，模态关系漂移约束以张量投影输出为固定参照，只惩罚领域转换造成的跨模态样本关系恶化。当前工作在So2Sat LCZ42 v4上采用跨城市监督域适应协议，"
        "以检验该框架处理区域分布差异的能力，并为后续跨季节、跨时相和跨传感器验证提供统一方法基础。",
    )
    add(document, "关键词：卫星遥感；多模态学习；监督域适应；SAR与光学融合；分布漂移；双向对抗学习")
    add(document, "中图分类号：TP391.41")
    add(document, "DOI：")

    add(document, "引言", "Heading 1")
    add(
        document,
        "卫星遥感模型通常在有限区域、时间窗口和观测条件下训练，却需要服务于持续变化的地表环境。太阳高度角、云雾和气溶胶会改变光学反射特征，"
        "季节物候与土地利用演化会改变类别内部结构，不同轨道、入射角和地表含水量则会引起SAR后向散射变化。即便任务的类别定义保持不变，"
        "上述因素仍会使训练数据与部署数据之间出现显著分布差异。地区变化进一步叠加城市形态、建筑材料、植被覆盖和地形条件的差异，"
        "使卫星分类系统面对的并非单一、静态的域偏移，而是由多种观测与地理因素共同驱动的动态分布漂移。",
    )
    add(
        document,
        "SAR与多光谱光学数据提供了互补的地表观测。光学影像包含连续的光谱响应和清晰的空间纹理，适合区分植被、裸地、建筑材料及水体；SAR主动成像对太阳光照依赖较弱，"
        "能够表征介电性质、表面粗糙度和几何结构，但同时受到相干斑、叠掩和阴影影响。两种模态的互补性使联合学习具有优势，也带来成像机理、通道统计和特征语义不一致的问题。"
        "简单拼接只能合并数值表示，不能解释一种模态如何调节另一种模态的判别关系；若在融合后直接施加强域对齐，模型还可能通过改变模态间结构来降低域差异。",
    )
    add(
        document,
        "遥感域适应方法主要通过统计矩匹配、对抗学习或类别条件约束缩小源域与目标域差异。多模态方法则通常采用中间融合、交叉注意力或共享—私有表示分解。"
        "这两条技术路线在许多工作中仍按串行方式组合：先获得融合表示，再对融合分布进行整体对齐。该策略忽略了两个事实。其一，跨域偏移会分别作用于SAR和光学模态，"
        "同一类别在两域中的对应关系应在多模态上下文中建模；其二，整体特征变换具有跨模态混合能力，后续对抗优化可能削弱前端建立的互补结构。"
        "因此，可靠的卫星多模态域适应不仅需要降低领域差异，还需要保持类别语义和模态关系。",
    )
    add(
        document,
        "基于这一认识，本文提出Satellite-Dual_D。方法按“建立多模态关系、学习领域转换、限制语义偏移、保护模态结构”的顺序组织。"
        "张量模块不直接要求SAR与光学特征数值相同，而是利用两种模态的联合响应调节各自表示，并对齐源域和目标域中同一模态的上下文增强特征。"
        "双向生成器随后学习难以由线性投影覆盖的非线性域差异；类别感知反馈约束转换后的类别边界；关系漂移项则检查全局转换是否破坏张量阶段形成的多模态关系。",
    )
    add(document, "本文的主要贡献如下：")
    add(
        document,
        "（1）提出一个面向卫星多模态分布漂移的双向域适应框架，将SAR—光学关系建模、跨域特征转换、类别保持和关系保护置于同一优化链路中。"
        "该框架适用于由地区、光照、季节和观测条件变化引起的源—目标分布差异；当前实验以跨城市偏移作为可复现验证场景。",
    )
    add(
        document,
        "（2）设计计算可行的张量式多模态关系建模与跨域对齐模块。该模块利用秩一外积的因子化收缩避免显式构造高维张量，"
        "在多模态上下文中对齐源域和目标域的同模态表示，并以域特定正交投影形成下游融合特征。",
    )
    add(
        document,
        "（3）构建双向残差生成器、方向特定双判别器和多尺度类别感知反馈，并引入单侧模态关系漂移约束。漂移项以停止梯度的张量投影结果为基线，"
        "允许领域转换所需的局部变化，只限制超过容忍边界的关系恶化。",
    )
    add(document, "本文其余部分组织如下：第二节回顾卫星多模态学习与域适应研究；第三节给出Satellite-Dual_D的完整方法；第四节说明数据、协议和实验设置；第五节总结本文。")

    add(document, "相关工作", "Heading 1")
    add(document, "卫星多模态表征学习", "Heading 2")
    add(
        document,
        "多源卫星观测能够从光谱、散射、几何和时间变化等不同角度描述地表。以So2Sat LCZ42为代表的数据集提供同一地理斑块的Sentinel-1与Sentinel-2观测，"
        "推动了SAR—光学联合分类研究 \\cite{schmitt2019so2sat}。现有方法多采用特征拼接、决策级融合、门控加权或跨模态注意力建立联合表示。"
        "这些方法能够利用模态互补性，但通常假设训练与测试数据同分布；当城市形态、地表材料或采集条件变化时，融合层可能放大模态间不稳定相关性。",
    )
    add(
        document,
        "另一类研究通过共享—私有特征分解或对比学习区分模态共有语义与模态特有信息。该思路有助于避免强制同质化，但多数约束仍集中在同一数据域内部。"
        "对于跨地区、跨季节或跨传感器任务，模态关系本身也会随领域变化，因此需要同时处理模态异质性和跨域偏移。",
    )
    add(document, "卫星遥感中的域适应", "Heading 2")
    add(
        document,
        "遥感域适应通常通过最大均值差异、协方差匹配或对抗域判别学习域不变表示 \\cite{long2015learning,sun2016deep,ganin2016domain}。"
        "类别条件方法进一步利用原型、伪标签或类内紧致性缓解全局对齐造成的类别混淆。对卫星数据而言，域差异往往由地区与观测条件共同构成，"
        "仅匹配整体边缘分布可能把具有相似低层统计、但语义不同的地表类型拉近。",
    )
    add(
        document,
        "当前Satellite-Dual_D采用监督式跨区域域适应：源域与目标适配集均提供类别标签，类别空间保持一致。"
        "这一设定不同于无监督域适应，也不等同于未见城市域泛化。目标标签用于类平衡跨域采样、目标域分类、同类对比、目标域类别原型和生成特征反馈，"
        "因而评价时必须同时报告目标域直接监督基线，才能区分复杂适配模块与普通联合监督训练的贡献。",
    )
    add(document, "多模态域对齐与结构保持", "Heading 2")
    add(
        document,
        "多模态域适应需要在模态交互和领域不变性之间取得平衡。“先交互、后对齐”的方法利用跨模态信息补充单一模态中缺失的可迁移线索 \\cite{Yang_2022_CVPR}；"
        "共享—私有分解与监督对比学习则从语义层面约束跨域表示 \\cite{NEURIPS2023_f88bec15}。OSAN进一步以高阶张量关系统一多模态交互与域对齐 \\cite{liu2023osan}。"
        "本文借鉴其高阶关系建模思想，但当前实现采用端到端因子化张量收缩和配对余弦相关性，不复现OSAN的交叉协方差白化与交替广义特征值求解。",
    )
    add(
        document,
        "现有研究较少显式检查后续领域变换是否损伤前端建立的模态协同关系。循环一致性和恒等约束控制整体特征变化，却不能判断SAR与光学对批内样本关系的描述是否同步退化。"
        "Satellite-Dual_D因此把关系保护作为对领域转换的附加约束，使域匹配不能以显著破坏预定义多模态子空间为代价。",
    )

    add(document, "方法", "Heading 1")
    add(document, "问题定义与总体框架", "Heading 2")
    add(
        document,
        "考虑包含SAR和多光谱光学观测的闭集卫星分类任务。源域数据来自一个地区或观测条件集合，目标域数据来自另一地区或条件集合，二者共享K个类别。"
        "训练集记为：",
    )
    add(document, "𝒟_d = {((x_(d,i)^sar, x_(d,i)^opt), y_(d,i))}_(i=1)^(N_d),   d ∈ {s,t}.    (1)", "Equation")
    add(
        document,
        "其中，P_s(x,y)≠P_t(x,y)，但类别空间满足𝒴_s=𝒴_t。单个样本内部的SAR与光学斑块对应同一地理位置；源域样本与目标域样本之间不要求空间或实例配对。"
        "当前实现使用目标适配集标签，并由类平衡采样器在同一批次的对应位置抽取同类跨域样本，即y_(s,i)=y_(t,i)。因此，本文任务是监督式跨区域域适应，"
        "不是无监督适配或在线连续适配。更一般的光照、季节和传感器漂移可用同一源—目标形式建模，但需要独立数据协议验证。",
    )
    add(
        document,
        "Satellite-Dual_D由五个环节组成。独立编码器首先分离处理SAR散射统计和光学光谱统计；张量模块在多模态上下文中对齐两域的同模态表示；"
        "双向特征翻译与两个方向特定判别器进一步匹配非线性域分布；类别感知反馈限制生成特征跨越类别边界；模态关系漂移约束则保护张量阶段建立的关系结构。"
        "这条链路把“降低领域差异”和“避免破坏已有结构”同时纳入优化。",
    )

    add(document, "模态特定卫星表征", "Heading 2")
    add(
        document,
        "输入由8通道Sentinel-1 SAR斑块x^sar∈ℝ^(8×32×32)和10通道Sentinel-2光学斑块x^opt∈ℝ^(10×32×32)组成。"
        "SAR通道包括VH、VV复数分量、滤波强度和协方差分量；光学通道覆盖可见光、红边、近红外和短波红外波段。为避免在底层统计上强制共享，"
        "本文为两种模态设置结构相同但参数独立的ResNet20-style编码器：",
    )
    add(document, "h_d^sar = E_sar(x_d^sar),   h_d^opt = E_opt(x_d^opt),   h_d^m ∈ ℝ^256.    (2)", "Equation")
    add(
        document,
        "E_sar与E_opt不共享参数，但同一模态在源域和目标域之间共享编码器。前者避免把两种成像机理混入同一低层卷积核，后者保证域适应针对观测分布变化，而不是为每个领域单独建立表征系统。"
        "每个编码器由三组残差阶段和全局平均池化构成，输出256维模态特征。",
    )

    add(document, "张量式多模态关系建模与跨域对齐", "Heading 2")
    add(
        document,
        "不同模态在本模块中承担两种不同作用：SAR与光学通过高阶响应相互调节，源域和目标域中的同一模态则接受显式相关性对齐。"
        "对模态m∈{sar,opt}，分别学习源域投影U_m∈ℝ^(256×128)与目标域投影V_m∈ℝ^(256×128)。投影矩阵采用正交初始化，"
        "并在每次主网络更新后通过QR分解回缩到列正交空间。",
    )
    add(
        document,
        "对样本i，联合响应可写为秩一外积𝒳_(d,i)=h_(d,i)^sar∘h_(d,i)^opt。显式构造外积会产生二次存储开销。利用秩一张量的可分解性，"
        "排除模态n后对另一模态进行投影和求和，可等价得到样本级门控与上下文增强表示：",
    )
    add(document, "a_(d,i)^n = ∏_(m≠n) 1^T(h_(d,i)^m W_(d,m)),   c_(d,i)^n = a_(d,i)^n h_(d,i)^n.    (3)", "Equation")
    add(
        document,
        "其中，W_(s,m)=U_m，W_(t,m)=V_m。a_(d,i)^n汇集另一模态的投影响应，使第n个模态的跨域比较受到多模态联合状态调节。"
        "由类平衡采样器产生的同类跨域样本在批次位置上对应，张量对齐损失为：",
    )
    add(document, "𝓛_TAL = −(1/2B) ∑_(n∈{sar,opt}) ∑_(i=1)^B cos(c_(s,i)^n, c_(t,i)^n).    (4)", "Equation")
    add(
        document,
        "式（4）提高的是同类跨域样本在同一模态下的上下文相关性，不直接要求h^sar与h^opt数值一致，也不能据此声称两域总体分布完全重合。"
        "用于下游任务的表示并非显式外积张量，而是域特定低维投影的拼接：",
    )
    add(document, "f_s = [h_s^sar U_sar ; h_s^opt U_opt],   f_t = [h_t^sar V_sar ; h_t^opt V_opt] ∈ ℝ^256.    (5)", "Equation")
    add(
        document,
        "因此，高阶交互通过𝓛_TAL间接塑造投影矩阵和编码器，分类器与后续翻译器接收的是两个128维模态投影块的融合向量。"
        "这一实现保留了固定的模态子空间边界，同时避免完整张量嵌入带来的参数与显存开销。",
    )

    add(document, "双向领域特征翻译与对抗分布匹配", "Heading 2")
    add(
        document,
        "相关性对齐难以覆盖由地区形态、成像几何和传感器响应共同造成的非线性差异。本文在256维融合空间中设置G_(s→t)和G_(t→s)，"
        "分别生成目标域化特征f_hat_t和源域化特征f_hat_s。两个生成器均采用残差多层感知机：",
    )
    add(document, "G_(a→b)(f) = f + ρ tanh(g_(a→b)(f)),   f_hat_t=G_(s→t)(f_s),   f_hat_s=G_(t→s)(f_t).    (6)", "Equation")
    add(
        document,
        "生成器作用于完整融合向量，能够同时调整两个预定义模态块并跨块混合信息；因此不能把该过程描述为分别翻译SAR和光学特征。"
        "源域判别器D_s区分f_s与f_hat_s，目标域判别器D_t区分f_t与f_hat_t。令ℓ_bin为二类交叉熵，判别器损失为：",
    )
    add(document, "𝓛_D^s=1/2[ℓ_bin(D_s(f_s),1)+ℓ_bin(D_s(f_hat_s),0)],", "Equation")
    add(document, "𝓛_D^t=1/2[ℓ_bin(D_t(f_t),1)+ℓ_bin(D_t(f_hat_t),0)],   𝓛_D=1/2(𝓛_D^s+𝓛_D^t).    (7)", "Equation")
    add(
        document,
        "生成器对抗目标为𝓛_adv^s=ℓ_bin(D_s(f_hat_s),1)和𝓛_adv^t=ℓ_bin(D_t(f_hat_t),1)。它们分别缩小生成源域分布与真实源域分布、"
        "生成目标域分布与真实目标域分布之间的差异，但不构成P_s=P_t的数学保证。两个方向的独立判别信号避免将不同域简单压入一个不可解释的混合分布。",
    )

    add(document, "类别感知语义反馈", "Heading 2")
    add(
        document,
        "域判别器只判断特征是否具有目标领域的统计特征，不知道转换前后的地表类别是否一致。本文以循环一致性和恒等保持限制信息破坏，以样本级对比、"
        "类别原型和分类器反馈约束类别语义。双向循环损失和恒等损失分别为：",
    )
    add(document, "𝓛_cyc=‖G_(t→s)(G_(s→t)(f_s))−f_s‖_1 + ‖G_(s→t)(G_(t→s)(f_t))−f_t‖_1,    (8)", "Equation")
    add(document, "𝓛_id=‖G_(t→s)(f_s)−f_s‖_1 + ‖G_(s→t)(f_t)−f_t‖_1.    (9)", "Equation")
    add(
        document,
        "循环一致性约束往返映射的可恢复性，恒等损失抑制生成器对已经位于输出域中的真实特征进行无意义改写。"
        "为显式保持类别结构，对生成锚点q_i、输出域真实候选r_j及同类集合P(i)={j|y_j=y_i}，定义同类多正样本对比损失：",
    )
    add(document, "𝓛_mp(q,r)=−(1/B)∑_i log[∑_(j∈P(i)) exp(sim(q_i,r_j)/τ) / ∑_j exp(sim(q_i,r_j)/τ)].    (10)", "Equation")
    add(
        document,
        "模型对(f_hat_s,f_s)与(f_hat_t,f_t)两个方向分别计算式（10）并取均值。全部同类候选共同构成正样本概率质量，异类候选作为负样本；"
        "这里的配对依据是类别标签，而不是同一地理斑块的跨域实例对应。真实候选在该损失中停止梯度，使对比反馈主要更新生成路径。",
    )
    add(
        document,
        "批内样本参照会随批次组成波动。为获得跨批次的类别尺度约束，本文分别维护源域和目标域的指数滑动原型。对类别c在当前批次中的均值p_bar_(d,c)，"
        "已有原型按p_(d,c)←μp_(d,c)+(1−μ)p_bar_(d,c)更新。f_hat_s与源域原型比较，f_hat_t与目标域原型比较，并以输入样本原标签作为交叉熵目标，得到𝓛_proto。"
        "尚未观测到的类别不参与原型损失。样本级对比提供当前批次的细粒度参照，原型对比提供跨批次的稳定类别参照。",
    )
    add(
        document,
        "共享分类器C进一步检查生成特征的可判别性：𝓛_fb=CE(C(f_hat_s),y_t)+CE(C(f_hat_t),y_s)。计算该项时冻结分类器参数，"
        "梯度仅通过C传回生成器，避免分类器迁就错误转换。真实融合特征的监督损失为𝓛_cls=CE(C(f_s),y_s)+λ_t CE(C(f_t),y_t)。",
    )

    add(document, "模态关系漂移约束", "Heading 2")
    add(
        document,
        "张量模块在两个128维投影块上建立多模态上下文关系，后续生成器却对完整256维向量进行非线性变换。即使类别保持正确，"
        "生成器仍可能通过改变两个子空间对样本结构的描述来降低域判别损失。漂移模块不执行新的领域对齐，而是检查这种关系是否相对张量输出显著恶化。",
    )
    add(
        document,
        "将融合特征f按固定边界切分为z^sar和z^opt。由于两种模态的坐标语义不同，本文不直接最小化z^sar与z^opt的欧氏距离，"
        "而分别构造批内余弦关系矩阵：",
    )
    add(document, "K^m = norm(z^m) norm(z^m)^T,   A(f)=(1/B^2)‖K^sar−K^opt‖_F^2.    (11)", "Equation")
    add(
        document,
        "A(f)衡量两个预定义模态子空间对批内样本相似关系的分歧。它不单独证明模态已经对齐，只用于比较领域转换前后的相对变化。"
        "考虑到生成器需要改变特征以完成域转换，本文设置容忍边界δ，并只惩罚关系不一致度的正向超额增量：",
    )
    add(document, "𝓛_drift=1/2([A(G_(s→t)(sg(f_s)))−sg(A(f_s))−δ]_+", "Equation")
    add(document, "+ [A(G_(t→s)(sg(f_t)))−sg(A(f_t))−δ]_+).    (12)", "Equation")
    add(
        document,
        "其中，sg表示停止梯度，[u]_+=max(u,0)。当转换后两种子空间的关系更一致，或恶化幅度不超过δ时，漂移项为零；"
        "只有超过阈值的恶化才产生梯度。参照特征停止梯度，因此该分支主要约束生成器，不会推动编码器和张量投影反向移动以迎合生成结果。"
        "漂移权重在训练后段逐步启用，避免在生成器尚未形成基本域转换能力时过早限制必要变化。由于生成器能够跨块混合信息，"
        "转换后的两段更严格地说是由张量投影边界定义的SAR与光学子空间位置，而非物理上完全独立的单模态特征。",
    )

    add(document, "总体目标与优化", "Heading 2")
    add(document, "𝓛_C=λ_id𝓛_id+λ_cyc𝓛_cyc+λ_mp𝓛_mp+λ_proto𝓛_proto+λ_fb𝓛_fb,    (13)", "Equation")
    add(document, "𝓛_adv=λ_s𝓛_adv^s+λ_t𝓛_adv^t,    (14)", "Equation")
    add(document, "𝓛_main=𝓛_cls+λ_TAL𝓛_TAL+γ_adv𝓛_adv+γ_C𝓛_C+γ_drift λ_drift𝓛_drift.    (15)", "Equation")
    add(
        document,
        "训练采用判别器与主网络交替更新。固定编码器、张量投影和生成器时，D_s与D_t最小化𝓛_D；固定判别器时，编码器、张量投影、生成器和分类器最小化𝓛_main。"
        "γ_adv、γ_C和γ_drift分别控制对抗分支、类别反馈和关系漂移的渐进启用。每次主网络更新后对U_m和V_m执行QR回缩。"
        "推理时只输入目标域SAR—光学样本，经目标域投影得到f_t，再由G_(t→s)生成源域化特征f_hat_s并送入分类器；判别器、反向生成器和辅助损失均不参与推理。",
    )

    experiment = add(document, "实验设置", "Heading 1")
    experiment.paragraph_format.page_break_before = True
    add(document, "数据集与跨区域协议", "Heading 2")
    add(
        document,
        "当前实验采用So2Sat LCZ42 v4 \\cite{schmitt2019so2sat}。数据集以32×32地理斑块为基本单元，每个样本包含8通道Sentinel-1 SAR和10通道Sentinel-2光学数据，"
        "类别空间由17类局地气候区组成。training.h5包含42个源城市的352,366个样本，作为源域；validation.h5包含10个目标城市的24,119个样本，作为目标适配域；"
        "testing.h5包含相同10个目标城市的独立空间分区，共24,188个样本，仅用于最终测试。",
    )
    add(
        document,
        "目标适配域内部按类别划分训练子集与检查点验证子集，划分比例为9:1，随机种子固定为2026。testing.h5不参与训练、超参数选择、早停或模型选择。"
        "所有归一化统计仅由源域训练数据估计，随后同样应用于目标适配集和测试集，避免利用目标测试分布计算预处理参数。"
        "随机分层划分可能保留相邻地理斑块之间的空间相关性，因此该协议用于当前跨城市监督适配实验；基于城市或空间块的更严格验证在结论中作为后续研究方向讨论。",
    )

    add(document, "实现细节", "Heading 2")
    add(
        document,
        "两个ResNet20-style编码器各输出256维特征，张量投影维度均为128，融合向量维度为256。生成器包含两层192维隐藏层，采用层归一化、ReLU、0.1 dropout和ρ=0.3的残差缩放。"
        "两个判别器的隐藏维度为128和64，使用谱归一化及0.5 dropout。类别原型动量μ=0.9，对比温度τ=0.2，关系漂移阈值δ=0.01。"
        "训练80个epoch，批大小为256；主网络与编码器学习率为3×10^−4，判别器学习率为1.5×10^−4，权重衰减为1×10^−4，梯度裁剪阈值为5。",
    )
    add(
        document,
        "损失权重设置为λ_TAL=0.3、λ_s=λ_t=0.035、λ_cyc=0.32、λ_id=0.05、λ_mp=0.04、λ_proto=0.06、λ_fb=0.5和λ_drift=0.02。"
        "对抗分支在前5个epoch关闭，随后用20个epoch线性增加至完整权重；类别反馈在前5个epoch关闭，并在随后10个epoch线性增加；漂移约束在前12个epoch仅记录，"
        "随后用12个epoch线性启用。模型选择基于目标适配验证子集的宏平均F1，并限制最早选择epoch为24。",
    )

    add(document, "对比与消融设置", "Heading 2")
    add(
        document,
        "为区分目标标签监督、张量关系建模和双向对抗模块的贡献，对比设置包括：仅使用源域训练的Source-only、仅使用目标适配标签的Target-only、"
        "源域与目标域直接联合监督、Deep CORAL、DANN、仅保留张量模块的TAL-only，以及完整Satellite-Dual_D。"
        "消融实验分别移除双向对抗分支、循环一致性、恒等保持、同类多正样本对比、EMA原型反馈、生成特征分类反馈和关系漂移约束。"
        "对双向翻译模块还应在同一检查点比较raw、source-like和残差融合三种推理路径，避免把分类器本身的监督收益误归因于特征翻译。",
    )
    add(
        document,
        "主要指标采用17类总体准确率与宏平均F1，并同时报告每类F1和混淆矩阵。所有主结果使用不少于三个随机种子，给出均值与标准差。"
        "关系漂移部分除最终分类指标外，还应报告两个方向的原始漂移量、超阈值比例及翻译前后的A(f)，从而验证该约束是否确实保护多模态子空间关系。",
    )

    add(document, "结论", "Heading 1")
    add(
        document,
        "本文面向卫星数据随光照、季节、地区、观测几何和传感器条件变化产生的多模态分布漂移，提出SAR—光学双对抗域适应网络Satellite-Dual_D。"
        "独立编码器与因子化张量收缩用于建立多模态上下文关系并对齐两域同模态表示，双向残差生成器和方向特定判别器进一步学习非线性领域转换。"
        "类别感知反馈限制语义偏移，模态关系漂移项以张量输出为固定基线，防止后续映射显著削弱预定义SAR与光学子空间的关系结构。"
        "当前So2Sat LCZ42协议验证的是跨城市监督域适应，不能直接外推为无监督、跨季节或在线动态适配。",
    )
def main():
    document = Document(SOURCE)
    clear_body(document)
    configure_styles(document)
    build_document(document)
    document.core_properties.title = "面向卫星多模态分布漂移的SAR与光学双对抗域适应网络"
    document.core_properties.subject = "Satellite-Dual_D最新方法与卫星遥感背景修订稿"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
